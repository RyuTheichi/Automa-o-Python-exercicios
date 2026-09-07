# Projeto – Relatório Acadêmico Automatizado

# Seu desafio é automatizar a geração de relatórios com base no desempenho dos alunos!

# Você receberá uma planilha chamada alunos.xlsx, contendo os dados de 30 alunos com as seguintes colunas:

#     Nome

#     Curso

#     Idade

#     Nota Final

#     Data de Matrícula

# 🧩 O que seu programa deve fazer:

#     Abrir a planilha

#     Percorrer todos os registros, separando os alunos em dois grupos:

#         Aprovados (nota final >= 7.0)

#         Reprovados (nota final < 7.0)

#     Criar dois novos arquivos Excel:

#         aprovados.xlsx

#         reprovados.xlsx

#     Em cada arquivo, salvar os dados completos dos respectivos alunos

#     Exibir no terminal:

#         Quantidade de aprovados e reprovados

#         Nota média da turma

#         Nome do aluno com a maior nota

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Color, Font, Alignment

wb_carregado = load_workbook(
    r"\Python_automation_class\manipulando_planilhas\Projeto\alunos.xlsx"
)
planilha_dados = wb_carregado["Alunos"]
qntd_aprovados = 0
qntd_reprovados = 0
soma = 0.0
maior_nota = 0.0

wb_aprovados = Workbook()
resultado_aprovado = wb_aprovados.active
resultado_aprovado.title = "Alunos aprovados"

wb_reprovados = Workbook()
resultado_reprovado = wb_reprovados.active
resultado_reprovado.title = "Alunos reprovados"

for linha in planilha_dados.iter_rows():
    if linha[0].row == 1:
        continue

    nota_celula = linha[3].value
    nota = float(nota_celula)
    soma = nota + soma
    if nota > maior_nota:
        maior_nota = nota
    valores_linha = [celula.value for celula in linha]
    if nota >= 7:
        resultado_aprovado.append(valores_linha)
        qntd_aprovados += 1
    else:
        resultado_reprovado.append(valores_linha)
        qntd_reprovados += 1
total_notas = qntd_reprovados + qntd_aprovados

print(f"Foram {qntd_aprovados} aprovados e {qntd_reprovados} reprovados\n")
print(f"A media da turma foi {soma / total_notas}\n")
print(f"A maior nota da turma foi {maior_nota}")

wb_aprovados.save(
    r"\Python_automation_class\manipulando_planilhas\Projeto\Aprovados.xlsx"
)
wb_reprovados.save(
    r"\Python_automation_class\manipulando_planilhas\Projeto\Reprovados.xlsx"
)
