from openpyxl import load_workbook

wk = load_workbook("planilha_funcionarios.xlsx")
planilha = wk["Funcionários"]

for linha in planilha.iter_rows(values_only=True):
    print("-" * 30)
    for celula in linha:
        print(celula)
