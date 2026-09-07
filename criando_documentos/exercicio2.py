# Exercício 2 — Gerar uma tabela_valor de funcionários

# Crie um arquivo Word chamado funcionarios.docx com:

#     Um título com o texto "Lista de Funcionários".

#     Uma tabela_valor com 3 colunas: Nome, Cargo, Salário.

#     Preencha os dados com informações fictícias, por exemplo:

#     +--------------+----------------+---------+

#     | Nome         | Cargo           | Salário |

#     +--------------+----------------+---------+

#     | João Silva   | Analista        | 5000    |

#     | Maria Souza  | Desenvolvedora  | 6000    |

#     | Pedro Santos | Designer        | 4000    |

#     | Ana Lima     | Gerente         | 7000    |

#     | Lucas Costa  | Estagiário      | 1500    |

#     +--------------+----------------+---------+


#     Aplique um estilo na tabela_valor.

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

funcionarios = [
    ("Joao Silva", "Analista", "5000"),
    ("Maria Souza", "Desenvolvedora", "6000"),
    ("Pedro Santos", "Designer", "4000"),
    ("Ana Lima", "Gerente", "7000"),
    ("Lucas Costa", "Estagiário", "1500"),
]
documento = Document()

titulo = documento.add_heading("", level=1)
titulo.add_run("Lista de Funcionários").font.size = Pt(16)
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

tabela = documento.add_table(0, 3, "Light List Accent 1")

cabecalho = tabela.add_row().cells
cabecalho[0].text = "NOME"
cabecalho[1].text = "CARGO"
cabecalho[2].text = "SALARIO"

for funcionario in funcionarios:
    tabela_valor = tabela.add_row().cells
    tabela_valor[0].text = funcionario[0]
    tabela_valor[1].text = funcionario[1]
    tabela_valor[2].text = funcionario[2]

documento.save("funcionarios.docx")
