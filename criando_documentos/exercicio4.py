# Enunciado de Exercício: Consumindo documentos com python-docx
# 📝 Exercício

# Abra o arquivo relatorio_inicial.docx e realize as seguintes modificações. Salve o resultado como relatorio_editado.docx.

#     Altere o título para "Relatório Financeiro - Revisado".

#     No parágrafo introdutório, adicione ao final do texto: " Os valores estão sujeitos a conferência."

#     Na tabela:

#         Altere a despesa de Fevereiro para 8500.

#         Adicione uma nova linha com o mês "Abril", receita 13000 e despesa 9000.

#     Adicione um parágrafo ao final do documento com o texto: "Relatório revisado em Python em [data de hoje]", onde a data de hoje é gerada automaticamente no código.

#     Salve o documento com o nome relatorio_editado.docx.

from docx import Document
from datetime import datetime

documento = Document("relatorio_inicial.docx")


documento.paragraphs[0].text = "Relatorio Financeiro - Revisado"

documento.paragraphs[1].add_run("Os valores estao sujeitos a conferencia.")

tabela = documento.tables[0]

tabela.cell(2, 2).text = "8500"

linha = tabela.add_row().cells

linha[0].text = "Abril"
linha[1].text = "13000"
linha[2].text = "9000"

hora_atual = datetime.now().strftime("%d/%m/%Y")

hora = documento.add_paragraph(f"Relatório revisado em Python em {hora_atual}")

documento.save("relatorio_editado.docx")
